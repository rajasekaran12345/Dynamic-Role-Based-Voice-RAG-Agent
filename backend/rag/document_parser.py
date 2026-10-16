import os
from typing import List, Dict, Any
from pypdf import PdfReader
import docx
from pptx import Presentation
import pandas as pd

def extract_text_from_pdf(file_path: str) -> str:
    text = ""
    reader = PdfReader(file_path)
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        text += f"\n--- Page {i+1} ---\n" + page_text
    return text.strip()

def extract_text_from_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)

def extract_text_from_pptx(file_path: str) -> str:
    prs = Presentation(file_path)
    text_runs = []
    for i, slide in enumerate(prs.slides):
        text_runs.append(f"\n--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_runs.append(shape.text.strip())
    return "\n".join(text_runs)

def extract_text_from_tabular(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".csv":
        df = pd.read_csv(file_path)
    elif ext in [".xlsx", ".xls"]:
        df = pd.read_excel(file_path)
    else:
        return ""
    
    # Format tabular as markdown/string summary
    rows = []
    rows.append(f"Columns: {', '.join(df.columns.astype(str))}")
    rows.append(f"Total Rows: {len(df)}")
    # Add head preview and string representation
    sample = df.head(100).to_string(index=False)
    rows.append(sample)
    return "\n".join(rows)

def extract_text_from_file(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext in [".pptx", ".ppt"]:
        return extract_text_from_pptx(file_path)
    elif ext in [".csv", ".xlsx", ".xls"]:
        return extract_text_from_tabular(file_path)
    elif ext in [".txt", ".md", ".json"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

def chunk_text(text: str, filename: str, chunk_size: int = 700, chunk_overlap: int = 100) -> List[Dict[str, Any]]:
    if not text.strip():
        return []

    words = text.split()
    chunks = []
    
    # Approximate chunking by word count
    words_per_chunk = chunk_size // 5  # approx 5 chars per word
    overlap_words = chunk_overlap // 5
    
    start = 0
    chunk_idx = 0
    while start < len(words):
        end = min(start + words_per_chunk, len(words))
        chunk_content = " ".join(words[start:end])
        if chunk_content.strip():
            chunks.append({
                "chunk_id": f"{filename}_{chunk_idx}",
                "chunk_index": chunk_idx,
                "filename": filename,
                "text": chunk_content
            })
            chunk_idx += 1
        if end >= len(words):
            break
        start += (words_per_chunk - overlap_words)
        
    return chunks
